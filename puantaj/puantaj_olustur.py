#!/usr/bin/env python3
"""
Yıllık Puantaj ve Mesai Takip Sistemi oluşturucu.

Kullanım:
    python puantaj_olustur.py                 # içinde bulunulan yıl
    python puantaj_olustur.py --yil 2027
    python puantaj_olustur.py --yil 2027 --cikti Dosya.xlsx --ornek-veri-yok

Oluşan çalışma kitabı:
    Yillik_Rapor      : Tüm ayları dinamik formüllerle toplayan icmal sayfası
    Aylik_Rapor       : Açılır listeden seçilen ayın personel özeti + 12 ay karşılaştırması
    Personel_Listesi  : Personel bilgileri tek yerden girilir, ay sayfaları buradan beslenir
    Ocak ... Aralık   : Aylık puantaj ve mesai giriş sayfaları

Durum kodları:  +  Geldi (tam gün)   x  Yarım gün   -  Gelmedi
"""

import argparse
import calendar
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------- #
# Ayarlar
# --------------------------------------------------------------------------- #
PERSONEL_KAPASITE = 15          # Personel listesindeki satır sayısı (boş satırlar dahil)
VARSAYILAN_CIKTI = "Puantaj_Takip_Sistemi.xlsx"

AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
         "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
GUN_KISA = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]

ORNEK_PERSONEL = [
    ("1001", "Ahmet Yılmaz", "Üretim"),
    ("1002", "Ayşe Demir", "Muhasebe"),
    ("1003", "Mehmet Kaya", "Lojistik"),
    ("1004", "Fatma Şahin", "İnsan Kaynakları"),
    ("1005", "Mustafa Çelik", "Üretim"),
    ("1006", "Zeynep Arslan", "Satış"),
    ("1007", "Emre Öztürk", "Bakım-Onarım"),
]

# --------------------------------------------------------------------------- #
# Stil paleti
# --------------------------------------------------------------------------- #
FONT_ADI = "Segoe UI"

LACIVERT = "1F3864"
CELIK_MAVI = "2E75B6"
ACIK_MAVI = "DDEBF7"
ZEBRA = "F3F7FB"
HAFTASONU = "E7E9EE"
KENAR_GRI = "B4C6E7"
BEYAZ = "FFFFFF"

F_BASLIK = Font(name=FONT_ADI, size=16, bold=True, color=BEYAZ)
F_ALT_BASLIK = Font(name=FONT_ADI, size=10, italic=True, color=LACIVERT)
F_TABLO_BASLIK = Font(name=FONT_ADI, size=10, bold=True, color=BEYAZ)
F_ALT_TABLO_BASLIK = Font(name=FONT_ADI, size=9, bold=True, color=BEYAZ)
F_NORMAL = Font(name=FONT_ADI, size=10)
F_KALIN = Font(name=FONT_ADI, size=10, bold=True, color=LACIVERT)
F_KPI_DEGER = Font(name=FONT_ADI, size=16, bold=True, color=LACIVERT)
F_KPI_ETIKET = Font(name=FONT_ADI, size=9, bold=True, color=CELIK_MAVI)

DOLGU_LACIVERT = PatternFill("solid", fgColor=LACIVERT)
DOLGU_CELIK = PatternFill("solid", fgColor=CELIK_MAVI)
DOLGU_ACIK_MAVI = PatternFill("solid", fgColor=ACIK_MAVI)
DOLGU_ZEBRA = PatternFill("solid", fgColor=ZEBRA)
DOLGU_HAFTASONU = PatternFill("solid", fgColor=HAFTASONU)
DOLGU_BEYAZ = PatternFill("solid", fgColor=BEYAZ)

INCE = Side(style="thin", color=KENAR_GRI)
ORTA = Side(style="medium", color=LACIVERT)
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)

ORTALI = Alignment(horizontal="center", vertical="center", wrap_text=True)
SOLA = Alignment(horizontal="left", vertical="center", indent=1)

# Durum kodları için koşullu biçimlendirme renkleri (dolgu, yazı)
DURUM_RENKLERI = {
    "+": ("C6EFCE", "006100"),   # Geldi     -> açık yeşil / koyu yeşil
    "-": ("FFC7CE", "9C0006"),   # Gelmedi   -> açık kırmızı / koyu kırmızı
    "x": ("FFEB9C", "9C5700"),   # Yarım gün -> açık sarı / koyu sarı
}

# --------------------------------------------------------------------------- #
# Yerleşim sabitleri
# --------------------------------------------------------------------------- #
AY_BASLIK_SATIR = 4          # Gün numaraları
AY_ALT_BASLIK_SATIR = 5      # Drm / Mesai
AY_ILK_VERI = 6
AY_SON_VERI = AY_ILK_VERI + PERSONEL_KAPASITE - 1
ILK_GUN_SUTUN = 4            # D

PL_ILK_VERI = 5              # Personel_Listesi ilk veri satırı

OZET_BASLIKLARI = [
    ("Tam Gün\n(+)", 9),
    ("Yarım Gün\n(x)", 9),
    ("Gelmediği\nGün (-)", 9),
    ("Fiili\nÇalışılan Gün", 11),
    ("Toplam\nMesai (Saat)", 11),
]


def drm_sutun(gun):
    return ILK_GUN_SUTUN + 2 * (gun - 1)


def mesai_sutun(gun):
    return drm_sutun(gun) + 1


def ozet_ilk_sutun(gun_sayisi):
    return ILK_GUN_SUTUN + 2 * gun_sayisi


def hucre_stil(c, font=F_NORMAL, fill=None, align=ORTALI, border=KENAR, fmt=None):
    c.font = font
    c.alignment = align
    c.border = border
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def birlesik_baslik(ws, satir, ilk_sutun, son_sutun, metin, font, fill, yukseklik=None):
    ws.merge_cells(start_row=satir, start_column=ilk_sutun, end_row=satir, end_column=son_sutun)
    c = ws.cell(row=satir, column=ilk_sutun, value=metin)
    c.font = font
    c.fill = fill
    c.alignment = ORTALI
    for col in range(ilk_sutun, son_sutun + 1):
        ws.cell(row=satir, column=col).fill = fill
    if yukseklik:
        ws.row_dimensions[satir].height = yukseklik


# --------------------------------------------------------------------------- #
# Personel_Listesi
# --------------------------------------------------------------------------- #
def personel_sayfasi(wb, yil):
    ws = wb.create_sheet("Personel_Listesi")
    ws.sheet_properties.tabColor = CELIK_MAVI
    ws.sheet_view.showGridLines = False

    birlesik_baslik(ws, 1, 1, 3, f"{yil} PERSONEL LİSTESİ", F_BASLIK, DOLGU_LACIVERT, 32)
    ws.merge_cells("A2:C2")
    ws["A2"] = ("Personel bilgilerini yalnızca bu sayfaya girin; tüm ay sayfaları ve "
                "yıllık rapor otomatik güncellenir. Ad Soyad BÜYÜK HARF yazılmalıdır.")
    ws["A2"].font = F_ALT_BASLIK
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    for i, baslik in enumerate(["Sicil No", "Ad Soyad", "Departman"], start=1):
        hucre_stil(ws.cell(row=4, column=i, value=baslik), F_TABLO_BASLIK, DOLGU_CELIK)
    ws.row_dimensions[4].height = 24

    for idx in range(PERSONEL_KAPASITE):
        r = PL_ILK_VERI + idx
        fill = DOLGU_ZEBRA if idx % 2 else DOLGU_BEYAZ
        kayit = ORNEK_PERSONEL[idx] if idx < len(ORNEK_PERSONEL) else (None, None, None)
        sicil, ad, dept = kayit
        hucre_stil(ws.cell(row=r, column=1, value=sicil), fill=fill, fmt="@")
        hucre_stil(ws.cell(row=r, column=2, value=ad.upper() if ad else None), fill=fill, align=SOLA)
        hucre_stil(ws.cell(row=r, column=3, value=dept), fill=fill)

    son = PL_ILK_VERI + PERSONEL_KAPASITE - 1
    dv = DataValidation(
        type="custom",
        formula1=f"EXACT(B{PL_ILK_VERI},UPPER(B{PL_ILK_VERI}))",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Büyük Harf Gerekli",
        error="Ad Soyad BÜYÜK HARF ile yazılmalıdır.",
    )
    ws.add_data_validation(dv)
    dv.add(f"B{PL_ILK_VERI}:B{son}")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 22
    ws.freeze_panes = f"A{PL_ILK_VERI}"
    return ws


# --------------------------------------------------------------------------- #
# Aylık puantaj sayfaları
# --------------------------------------------------------------------------- #
def ay_sayfasi(wb, yil, ay_no, ornek_veri):
    ad = AYLAR[ay_no - 1]
    gun_sayisi = calendar.monthrange(yil, ay_no)[1]
    ws = wb.create_sheet(ad)
    ws.sheet_properties.tabColor = LACIVERT
    ws.sheet_view.showGridLines = False

    son_gun_sutun = mesai_sutun(gun_sayisi)
    oz1 = ozet_ilk_sutun(gun_sayisi)
    son_sutun = oz1 + len(OZET_BASLIKLARI) - 1

    # Başlık ve açıklama
    birlesik_baslik(ws, 1, 1, son_sutun, f"{ad.upper()} {yil} — PERSONEL PUANTAJ VE MESAİ ÇİZELGESİ",
                    F_BASLIK, DOLGU_LACIVERT, 32)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    ws.cell(row=2, column=1, value=f"Takvim gün sayısı: {gun_sayisi}  |  Durum kodları  →").font = F_ALT_BASLIK
    ws.cell(row=2, column=1).alignment = SOLA
    lejant = [("+", "Geldi"), ("x", "Yarım Gün"), ("-", "Gelmedi")]
    col = ILK_GUN_SUTUN
    for kod, aciklama in lejant:
        dolgu, yazi = DURUM_RENKLERI[kod]
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 3)
        c = ws.cell(row=2, column=col, value=f"{kod}  {aciklama}")
        c.font = Font(name=FONT_ADI, size=9, bold=True, color=yazi)
        c.alignment = ORTALI
        for cc in range(col, col + 4):
            ws.cell(row=2, column=cc).fill = PatternFill("solid", fgColor=dolgu)
            ws.cell(row=2, column=cc).border = KENAR
        col += 5
    ws.row_dimensions[2].height = 20

    # Sabit başlıklar (A-C)
    for i, baslik in enumerate(["Sicil No", "Ad Soyad", "Departman"], start=1):
        ws.merge_cells(start_row=AY_BASLIK_SATIR, start_column=i, end_row=AY_ALT_BASLIK_SATIR, end_column=i)
        hucre_stil(ws.cell(row=AY_BASLIK_SATIR, column=i, value=baslik), F_TABLO_BASLIK, DOLGU_LACIVERT)
        hucre_stil(ws.cell(row=AY_ALT_BASLIK_SATIR, column=i), F_TABLO_BASLIK, DOLGU_LACIVERT)

    # Gün başlıkları
    haftasonu_sutunlari = set()
    for gun in range(1, gun_sayisi + 1):
        d, m = drm_sutun(gun), mesai_sutun(gun)
        hafta_gunu = dt.date(yil, ay_no, gun).weekday()
        haftasonu = hafta_gunu >= 5
        if haftasonu:
            haftasonu_sutunlari.update((d, m))
        ust_dolgu = DOLGU_CELIK if not haftasonu else PatternFill("solid", fgColor="5B6B85")
        ws.merge_cells(start_row=AY_BASLIK_SATIR, start_column=d, end_row=AY_BASLIK_SATIR, end_column=m)
        hucre_stil(ws.cell(row=AY_BASLIK_SATIR, column=d, value=f"{gun}\n{GUN_KISA[hafta_gunu]}"),
                   F_TABLO_BASLIK, ust_dolgu)
        hucre_stil(ws.cell(row=AY_BASLIK_SATIR, column=m), F_TABLO_BASLIK, ust_dolgu)
        hucre_stil(ws.cell(row=AY_ALT_BASLIK_SATIR, column=d, value="Drm"), F_ALT_TABLO_BASLIK, DOLGU_LACIVERT)
        hucre_stil(ws.cell(row=AY_ALT_BASLIK_SATIR, column=m, value="Mesai"), F_ALT_TABLO_BASLIK, DOLGU_LACIVERT)
        ws.column_dimensions[get_column_letter(d)].width = 4.6
        ws.column_dimensions[get_column_letter(m)].width = 5.8

    # Özet başlıkları
    for i, (baslik, genislik) in enumerate(OZET_BASLIKLARI):
        c = oz1 + i
        ws.merge_cells(start_row=AY_BASLIK_SATIR, start_column=c, end_row=AY_ALT_BASLIK_SATIR, end_column=c)
        hucre_stil(ws.cell(row=AY_BASLIK_SATIR, column=c, value=baslik), F_TABLO_BASLIK, DOLGU_LACIVERT)
        hucre_stil(ws.cell(row=AY_ALT_BASLIK_SATIR, column=c), F_TABLO_BASLIK, DOLGU_LACIVERT)
        ws.column_dimensions[get_column_letter(c)].width = genislik
    ws.row_dimensions[AY_BASLIK_SATIR].height = 30
    ws.row_dimensions[AY_ALT_BASLIK_SATIR].height = 18

    ilk_gun_harf = get_column_letter(ILK_GUN_SUTUN)
    son_gun_harf = get_column_letter(son_gun_sutun)
    oz = [get_column_letter(oz1 + i) for i in range(len(OZET_BASLIKLARI))]

    rng = random.Random(yil * 100 + ay_no)
    for idx in range(PERSONEL_KAPASITE):
        r = AY_ILK_VERI + idx
        pl = PL_ILK_VERI + idx
        zebra = DOLGU_ZEBRA if idx % 2 else DOLGU_BEYAZ

        # Personel bilgileri Personel_Listesi sayfasından gelir
        hucre_stil(ws.cell(row=r, column=1, value=f'=IF(Personel_Listesi!A{pl}="","",Personel_Listesi!A{pl})'),
                   fill=zebra)
        hucre_stil(ws.cell(row=r, column=2, value=f'=IF(Personel_Listesi!B{pl}="","",UPPER(Personel_Listesi!B{pl}))'),
                   F_KALIN, zebra, SOLA)
        hucre_stil(ws.cell(row=r, column=3, value=f'=IF(Personel_Listesi!C{pl}="","",Personel_Listesi!C{pl})'),
                   fill=zebra)

        ornek_satir = ornek_veri and idx < len(ORNEK_PERSONEL)
        for gun in range(1, gun_sayisi + 1):
            d, m = drm_sutun(gun), mesai_sutun(gun)
            fill = DOLGU_HAFTASONU if d in haftasonu_sutunlari else (zebra if idx % 2 else DOLGU_BEYAZ)
            drm_deger, mesai_deger = None, None
            if ornek_satir and d not in haftasonu_sutunlari:
                drm_deger = rng.choices(["+", "x", "-"], weights=[86, 8, 6])[0]
                if drm_deger == "+" and rng.random() < 0.3:
                    mesai_deger = rng.choice([1, 1.5, 2, 2.5, 3])
            hucre_stil(ws.cell(row=r, column=d, value=drm_deger), Font(name=FONT_ADI, size=10, bold=True),
                       fill, fmt="@")
            hucre_stil(ws.cell(row=r, column=m, value=mesai_deger), fill=fill, fmt="0.0;-0.0;;@")

        gun_araligi = f"{ilk_gun_harf}{r}:{son_gun_harf}{r}"
        mesai_hucreleri = ",".join(f"{get_column_letter(mesai_sutun(g))}{r}" for g in range(1, gun_sayisi + 1))
        formuller = [
            f'=IF($A{r}="","",COUNTIF({gun_araligi},"+"))',
            f'=IF($A{r}="","",COUNTIF({gun_araligi},"x"))',
            f'=IF($A{r}="","",COUNTIF({gun_araligi},"-"))',
            f'=IF($A{r}="","",{oz[0]}{r}+{oz[1]}{r}*0.5)',
            f'=IF($A{r}="","",SUM({mesai_hucreleri}))',
        ]
        bicimler = ["0", "0", "0", "0.0", "0.0"]
        for i, (formul, fmt) in enumerate(zip(formuller, bicimler)):
            hucre_stil(ws.cell(row=r, column=oz1 + i, value=formul), F_KALIN, DOLGU_ACIK_MAVI, fmt=fmt)
        ws.row_dimensions[r].height = 20

    # Özet bloğunu ayıran kalın kenarlık
    for r in range(AY_BASLIK_SATIR, AY_SON_VERI + 1):
        c = ws.cell(row=r, column=oz1)
        c.border = Border(left=ORTA, right=INCE, top=INCE, bottom=INCE)

    # Veri doğrulama: Drm -> liste, Mesai -> 0-24 arası sayı
    drm_aralik = " ".join(
        f"{get_column_letter(drm_sutun(g))}{AY_ILK_VERI}:{get_column_letter(drm_sutun(g))}{AY_SON_VERI}"
        for g in range(1, gun_sayisi + 1))
    mesai_aralik = " ".join(
        f"{get_column_letter(mesai_sutun(g))}{AY_ILK_VERI}:{get_column_letter(mesai_sutun(g))}{AY_SON_VERI}"
        for g in range(1, gun_sayisi + 1))
    dv_drm = DataValidation(type="list", formula1='"+,x,-"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Geçersiz Durum",
                            error="Lütfen + (Geldi), x (Yarım Gün) veya - (Gelmedi) girin.",
                            promptTitle="Durum", prompt="+ Geldi | x Yarım Gün | - Gelmedi",
                            showInputMessage=True)
    dv_mesai = DataValidation(type="decimal", operator="between", formula1="0", formula2="24",
                              allow_blank=True, showErrorMessage=True, errorTitle="Geçersiz Mesai",
                              error="Mesai 0 ile 24 arasında bir sayı (saat) olmalıdır.")
    ws.add_data_validation(dv_drm)
    ws.add_data_validation(dv_mesai)
    for aralik in drm_aralik.split():
        dv_drm.add(aralik)
    for aralik in mesai_aralik.split():
        dv_mesai.add(aralik)

    # Koşullu biçimlendirme (tüm gün alanı; Mesai sütunları sayısal olduğundan etkilenmez)
    cf_aralik = f"{ilk_gun_harf}{AY_ILK_VERI}:{son_gun_harf}{AY_SON_VERI}"
    for kod, (dolgu, yazi) in DURUM_RENKLERI.items():
        ws.conditional_formatting.add(
            cf_aralik,
            FormulaRule(formula=[f'{ilk_gun_harf}{AY_ILK_VERI}="{kod}"'],
                        fill=PatternFill("solid", fgColor=dolgu, bgColor=dolgu),
                        font=Font(bold=True, color=yazi),
                        stopIfTrue=True))

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 17
    ws.freeze_panes = ws.cell(row=AY_ILK_VERI, column=ILK_GUN_SUTUN)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_cols = "A:C"
    ws.print_title_rows = f"{AY_BASLIK_SATIR}:{AY_ALT_BASLIK_SATIR}"

    # Rapor sayfasının kullanacağı özet sütun harfleri
    return {"tam": oz[0], "yarim": oz[1], "gelmedi": oz[2], "fiili": oz[3], "mesai": oz[4],
            "gun_sayisi": gun_sayisi}


# --------------------------------------------------------------------------- #
# Yıllık rapor
# --------------------------------------------------------------------------- #
def rapor_sayfasi(ws, yil, ay_ozet_sutunlari):
    ws.title = "Yillik_Rapor"
    ws.sheet_properties.tabColor = "C00000"
    ws.sheet_view.showGridLines = False

    basliklar = ["Sicil No", "Ad Soyad", "Departman", "Ay",
                 "Toplam Çalışılan Gün", "Toplam Devamsızlık", "Toplam Mesai (Saat)"]
    genislikler = [12, 26, 20, 16, 18, 18, 18]
    n = len(basliklar)

    birlesik_baslik(ws, 1, 1, n, f"{yil} YILI PERSONEL PUANTAJ VE MESAİ İCMAL RAPORU",
                    F_BASLIK, DOLGU_LACIVERT, 36)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws["A2"] = ("Bu sayfadaki tüm değerler aylık puantaj sayfalarına bağlı formüllerdir; "
                "elle değer girmeyiniz. Devamsızlık = '-' (gelmediği gün) sayısı.")
    ws["A2"].font = F_ALT_BASLIK
    ws["A2"].alignment = ORTALI
    ws.row_dimensions[2].height = 20

    ilk = 9
    son = ilk + PERSONEL_KAPASITE * 13 - 1

    # KPI kartları
    kpi = [
        ("TOPLAM PERSONEL", f'=COUNTA(Personel_Listesi!A{PL_ILK_VERI}:A{PL_ILK_VERI + PERSONEL_KAPASITE - 1})', "0"),
        ("TOPLAM ÇALIŞILAN GÜN", f'=SUMIF($D${ilk}:$D${son},"YILLIK TOPLAM",$E${ilk}:$E${son})', "#,##0.0"),
        ("TOPLAM DEVAMSIZLIK", f'=SUMIF($D${ilk}:$D${son},"YILLIK TOPLAM",$F${ilk}:$F${son})', "#,##0"),
        ("TOPLAM MESAİ (SAAT)", f'=SUMIF($D${ilk}:$D${son},"YILLIK TOPLAM",$G${ilk}:$G${son})', "#,##0.0"),
    ]
    kpi_sutunlari = [(1, 2), (3, 4), (5, 5), (6, 7)]
    for (etiket, formul, fmt), (c1, c2) in zip(kpi, kpi_sutunlari):
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        e = ws.cell(row=4, column=c1, value=etiket)
        v = ws.cell(row=5, column=c1, value=formul)
        hucre_stil(e, F_KPI_ETIKET, DOLGU_ACIK_MAVI, border=Border())
        hucre_stil(v, F_KPI_DEGER, DOLGU_ACIK_MAVI, border=Border(), fmt=fmt)
        for cc in range(c1, c2 + 1):
            ws.cell(row=4, column=cc).fill = DOLGU_ACIK_MAVI
            ws.cell(row=4, column=cc).border = Border(top=ORTA)
            ws.cell(row=5, column=cc).fill = DOLGU_ACIK_MAVI
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 30

    # Tablo başlıkları
    for i, (baslik, genislik) in enumerate(zip(basliklar, genislikler), start=1):
        hucre_stil(ws.cell(row=7, column=i, value=baslik), F_TABLO_BASLIK, DOLGU_LACIVERT)
        ws.column_dimensions[get_column_letter(i)].width = genislik
    ws.row_dimensions[7].height = 28

    r = ilk - 1
    ws.row_dimensions[r].height = 4
    for idx in range(PERSONEL_KAPASITE):
        ay_r = AY_ILK_VERI + idx
        blok_ilk = r + 1
        for ay_no, ay in enumerate(AYLAR, start=1):
            r += 1
            s = ay_ozet_sutunlari[ay]
            ref = f"'{ay}'!"
            zebra = DOLGU_ZEBRA if ay_no % 2 == 0 else DOLGU_BEYAZ
            degerler = [
                f'={ref}$A${ay_r}',
                f'={ref}$B${ay_r}',
                f'={ref}$C${ay_r}',
                f'=IF({ref}$A${ay_r}="","","{ay}")',
                f'={ref}${s["fiili"]}${ay_r}',
                f'={ref}${s["gelmedi"]}${ay_r}',
                f'={ref}${s["mesai"]}${ay_r}',
            ]
            bicimler = [None, None, None, None, "0.0", "0", "0.0"]
            for c, (v, fmt) in enumerate(zip(degerler, bicimler), start=1):
                hucre_stil(ws.cell(row=r, column=c, value=v), fill=zebra,
                           align=SOLA if c == 2 else ORTALI, fmt=fmt)

        # Personel yıllık toplam satırı
        r += 1
        blok_son = r - 1
        a = f"$A${blok_ilk}"
        toplam = [
            f"={a}",
            f"=$B${blok_ilk}",
            f"=$C${blok_ilk}",
            f'=IF({a}="","","YILLIK TOPLAM")',
            f'=IF({a}="","",SUM(E{blok_ilk}:E{blok_son}))',
            f'=IF({a}="","",SUM(F{blok_ilk}:F{blok_son}))',
            f'=IF({a}="","",SUM(G{blok_ilk}:G{blok_son}))',
        ]
        bicimler = [None, None, None, None, "0.0", "0", "0.0"]
        for c, (v, fmt) in enumerate(zip(toplam, bicimler), start=1):
            hucre_stil(ws.cell(row=r, column=c, value=v), F_KALIN, DOLGU_ACIK_MAVI,
                       align=SOLA if c == 2 else ORTALI,
                       border=Border(left=INCE, right=INCE, top=INCE, bottom=ORTA), fmt=fmt)
        ws.row_dimensions[r].height = 20

    ws.auto_filter.ref = f"A7:{get_column_letter(n)}{son}"
    ws.freeze_panes = "A8"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "7:7"


# --------------------------------------------------------------------------- #
# Aylık rapor
# --------------------------------------------------------------------------- #
def aylik_rapor_sayfasi(wb, yil, ay_ozet_sutunlari):
    ws = wb.create_sheet("Aylik_Rapor", 1)
    ws.sheet_properties.tabColor = "C55A11"
    ws.sheet_view.showGridLines = False

    basliklar = ["Sicil No", "Ad Soyad", "Departman", "Tam Gün", "Yarım Gün", "Gelmediği Gün",
                 "Fiili Çalışılan Gün", "Toplam Mesai (Saat)", "Devam Oranı"]
    genislikler = [12, 26, 20, 11, 11, 13, 15, 15, 12]
    n = len(basliklar)
    for i, g in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(i)].width = g

    birlesik_baslik(ws, 1, 1, n, f"{yil} YILI AYLIK PERSONEL PUANTAJ VE MESAİ RAPORU",
                    F_BASLIK, DOLGU_LACIVERT, 36)

    # Ay seçici
    ws.merge_cells("A3:B3")
    ws["A3"] = "RAPOR AYI  ▸"
    hucre_stil(ws["A3"], F_TABLO_BASLIK, DOLGU_CELIK)
    hucre_stil(ws["B3"], F_TABLO_BASLIK, DOLGU_CELIK)
    secici = ws["C3"]
    secici.value = AYLAR[0]
    hucre_stil(secici, Font(name=FONT_ADI, size=12, bold=True, color="C00000"), PatternFill("solid", fgColor="FFF2CC"),
               border=Border(left=ORTA, right=ORTA, top=ORTA, bottom=ORTA))
    dv = DataValidation(type="list", formula1='"' + ",".join(AYLAR) + '"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Geçersiz Ay", error="Listeden bir ay seçiniz.",
                        showInputMessage=True, promptTitle="Rapor Ayı", prompt="Raporlanacak ayı listeden seçin.")
    ws.add_data_validation(dv)
    dv.add("C3")
    ws.merge_cells(start_row=3, start_column=4, end_row=3, end_column=n)
    ws["D3"] = '="Takvim gün sayısı: "&CHOOSE(MATCH($C$3,{' + ",".join(f'"{a}"' for a in AYLAR) + '},0),' + \
        ",".join(str(ay_ozet_sutunlari[a]["gun_sayisi"]) for a in AYLAR) + ')&"   |   Ay seçimini sarı hücreden değiştirin."'
    ws["D3"].font = F_ALT_BASLIK
    ws["D3"].alignment = SOLA
    ws.row_dimensions[3].height = 24

    ay_listesi = "{" + ",".join(f'"{a}"' for a in AYLAR) + "}"
    ay_idx = f"MATCH($C$3,{ay_listesi},0)"

    def secili_ay(anahtar, satir):
        refs = [f"'{a}'!${ay_ozet_sutunlari[a][anahtar]}${satir}" for a in AYLAR]
        return f"CHOOSE({ay_idx},{','.join(refs)})"

    # Seçili ay tablosu
    bas = 5
    for i, b in enumerate(basliklar, start=1):
        hucre_stil(ws.cell(row=bas, column=i, value=b), F_TABLO_BASLIK, DOLGU_LACIVERT)
    ws.row_dimensions[bas].height = 28
    ilk = bas + 1
    for idx in range(PERSONEL_KAPASITE):
        r = ilk + idx
        ay_r = AY_ILK_VERI + idx
        zebra = DOLGU_ZEBRA if idx % 2 else DOLGU_BEYAZ
        a = f"$A{r}"
        formuller = [
            f'=IF(Personel_Listesi!A{PL_ILK_VERI + idx}="","",Personel_Listesi!A{PL_ILK_VERI + idx})',
            f'=IF({a}="","",UPPER(Personel_Listesi!B{PL_ILK_VERI + idx}))',
            f'=IF({a}="","",Personel_Listesi!C{PL_ILK_VERI + idx})',
            f'=IF({a}="","",{secili_ay("tam", ay_r)})',
            f'=IF({a}="","",{secili_ay("yarim", ay_r)})',
            f'=IF({a}="","",{secili_ay("gelmedi", ay_r)})',
            f'=IF({a}="","",{secili_ay("fiili", ay_r)})',
            f'=IF({a}="","",{secili_ay("mesai", ay_r)})',
            f'=IF({a}="","",IF(D{r}+E{r}+F{r}=0,"",G{r}/(D{r}+E{r}+F{r})))',
        ]
        bicimler = ["@", None, None, "0", "0", "0", "0.0", "0.0", "0.0%"]
        for c, (f, fmt) in enumerate(zip(formuller, bicimler), start=1):
            hucre_stil(ws.cell(row=r, column=c, value=f), F_KALIN if c == 2 else F_NORMAL, zebra,
                       SOLA if c == 2 else ORTALI, fmt=fmt)
    son = ilk + PERSONEL_KAPASITE - 1
    t = son + 1
    hucre_stil(ws.cell(row=t, column=1, value="AY TOPLAMI"), F_TABLO_BASLIK, DOLGU_CELIK)
    ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=3)
    for c in range(2, 4):
        hucre_stil(ws.cell(row=t, column=c), F_TABLO_BASLIK, DOLGU_CELIK)
    for c, fmt in zip(range(4, 9), ["0", "0", "0", "0.0", "0.0"]):
        h = get_column_letter(c)
        hucre_stil(ws.cell(row=t, column=c, value=f"=SUM({h}{ilk}:{h}{son})"), F_TABLO_BASLIK, DOLGU_CELIK, fmt=fmt)
    hucre_stil(ws.cell(row=t, column=9, value=f'=IF(D{t}+E{t}+F{t}=0,"",G{t}/(D{t}+E{t}+F{t}))'),
               F_TABLO_BASLIK, DOLGU_CELIK, fmt="0.0%")
    ws.row_dimensions[t].height = 22

    # 12 aylık karşılaştırma (tüm personel)
    k = t + 3
    birlesik_baslik(ws, k, 1, n, "AYLARA GÖRE KARŞILAŞTIRMA (TÜM PERSONEL)", F_TABLO_BASLIK, DOLGU_LACIVERT, 24)
    kb = ["Ay", "", "Takvim Günü", "Tam Gün", "Yarım Gün", "Gelmediği Gün", "Fiili Çalışılan Gün",
          "Toplam Mesai (Saat)", "Devam Oranı"]
    for i, b in enumerate(kb, start=1):
        hucre_stil(ws.cell(row=k + 1, column=i, value=b or None), F_TABLO_BASLIK, DOLGU_CELIK)
    ws.merge_cells(start_row=k + 1, start_column=1, end_row=k + 1, end_column=2)
    ws.row_dimensions[k + 1].height = 28
    for ay_no, ay in enumerate(AYLAR, start=1):
        r = k + 1 + ay_no
        s = ay_ozet_sutunlari[ay]
        zebra = DOLGU_ZEBRA if ay_no % 2 == 0 else DOLGU_BEYAZ
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        hucre_stil(ws.cell(row=r, column=1, value=ay), F_KALIN, zebra)
        hucre_stil(ws.cell(row=r, column=2), fill=zebra)
        hucre_stil(ws.cell(row=r, column=3, value=s["gun_sayisi"]), fill=zebra)
        for c, anahtar, fmt in [(4, "tam", "0"), (5, "yarim", "0"), (6, "gelmedi", "0"),
                                (7, "fiili", "0.0"), (8, "mesai", "0.0")]:
            h = s[anahtar]
            hucre_stil(ws.cell(row=r, column=c, value=f"=SUM('{ay}'!{h}{AY_ILK_VERI}:{h}{AY_SON_VERI})"),
                       fill=zebra, fmt=fmt)
        hucre_stil(ws.cell(row=r, column=9, value=f'=IF(D{r}+E{r}+F{r}=0,"",G{r}/(D{r}+E{r}+F{r}))'),
                   fill=zebra, fmt="0.0%")
    # Seçili ayı karşılaştırma tablosunda vurgula
    ws.conditional_formatting.add(
        f"A{k + 2}:I{k + 13}",
        FormulaRule(formula=[f"$A{k + 2}=$C$3"], fill=PatternFill("solid", fgColor="FFF2CC", bgColor="FFF2CC"),
                    font=Font(bold=True, color="C00000")))
    yt = k + 14
    ws.merge_cells(start_row=yt, start_column=1, end_row=yt, end_column=2)
    hucre_stil(ws.cell(row=yt, column=1, value="YIL TOPLAMI"), F_TABLO_BASLIK, DOLGU_LACIVERT)
    hucre_stil(ws.cell(row=yt, column=2), F_TABLO_BASLIK, DOLGU_LACIVERT)
    for c, fmt in zip(range(3, 9), ["0", "0", "0", "0", "0.0", "0.0"]):
        h = get_column_letter(c)
        hucre_stil(ws.cell(row=yt, column=c, value=f"=SUM({h}{k + 2}:{h}{k + 13})"),
                   F_TABLO_BASLIK, DOLGU_LACIVERT, fmt=fmt)
    hucre_stil(ws.cell(row=yt, column=9, value=f'=IF(D{yt}+E{yt}+F{yt}=0,"",G{yt}/(D{yt}+E{yt}+F{yt}))'),
               F_TABLO_BASLIK, DOLGU_LACIVERT, fmt="0.0%")

    ws.freeze_panes = f"A{ilk}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --------------------------------------------------------------------------- #
def olustur(yil, cikti, ornek_veri=True):
    wb = Workbook()
    rapor_ws = wb.active                     # İlk sayfa: Yillik_Rapor
    personel_sayfasi(wb, yil)
    ozetler = {}
    for ay_no in range(1, 13):
        # Örnek puantaj verisi yalnızca Ocak ayına işlenir
        ozetler[AYLAR[ay_no - 1]] = ay_sayfasi(wb, yil, ay_no, ornek_veri and ay_no == 1)
    rapor_sayfasi(rapor_ws, yil, ozetler)
    aylik_rapor_sayfasi(wb, yil, ozetler)
    wb.active = 0
    wb.save(cikti)
    return cikti


def main():
    p = argparse.ArgumentParser(description="Yıllık Excel puantaj ve mesai takip sistemi oluşturur.")
    p.add_argument("--yil", type=int, default=dt.date.today().year, help="Puantaj yılı (varsayılan: bu yıl)")
    p.add_argument("--cikti", default=VARSAYILAN_CIKTI, help=f"Çıktı dosyası (varsayılan: {VARSAYILAN_CIKTI})")
    p.add_argument("--ornek-veri-yok", action="store_true",
                   help="Ocak ayına örnek puantaj verisi işlenmesin")
    args = p.parse_args()
    yol = olustur(args.yil, args.cikti, ornek_veri=not args.ornek_veri_yok)
    print(f"✔ {yol} oluşturuldu ({args.yil}).")


if __name__ == "__main__":
    main()
